#!/usr/bin/env python3
"""Build data/fin_catalogue.json from the AZM finance questionnaire spreadsheet.

The spreadsheet is the source of truth for the questionnaire: modules, sub-topics,
question wording, per-question options and the follow-up visibility conditions. It does
NOT carry the answer-shape metadata (title, profile, scope, table columns, branches),
which only exists in the catalogue already shipped with the addon. So this script
merges the two, keyed on the requirement code, and the spreadsheet wins wherever both
have a value.

It is a build-time tool: it is never imported by the addon, which keeps openpyxl out of
the runtime dependencies. Run it whenever a new revision of the sheet arrives:

    python3 tools/build_fin_catalogue.py AZM_Finance_Questions_Extract-2.xlsx

Nothing is written unless every check in validate() passes. Importing a questionnaire
where a question shows options belonging to a different question is worse than not
importing at all, so the checks are failures, not warnings.
"""

import argparse
import json
import pathlib
import re
import sys
from collections import Counter, defaultdict

import openpyxl

HERE = pathlib.Path(__file__).resolve().parent
ADDON = HERE.parent
CATALOGUE = ADDON / "data" / "fin_catalogue.json"

SHEET_QUESTIONS = "الأسئلة وطريقة الإجابة"
SHEET_OPENING = "الأسئلة الافتتاحية"

# Column headers, kept in one place so a renamed column fails loudly and early.
C_MODULE_CODE = "كود الموديول"
C_MODULE = "الموديول"
C_TOPIC_CODE = "كود الموضوع"
C_TOPIC = "الموضوع الفرعي"
C_REQ_ID = "Requirement ID"
C_REQ_CODE = "الكود السابق"
C_Q_CODE = "كود السؤال"
C_ROLE = "الدور"
C_QUESTION = "السؤال"
C_EXPECTED = "المطلوب في الإجابة"
C_FORMAT = "شكل الإجابة"
C_UNIT = "الوحدة"
C_COLUMNS = "أعمدة الجدول"
C_OPTIONS = "الاختيارات"
C_CONDITION = "شرط الظهور"
C_TEXTBOX = "خانة نص"
C_RESPONSE = "نوع الرد (V3)"
C_PACK = "حزمة الانطباق"
C_OWNER = "المالك المقترح"

C_NODE = "Node ID"
C_AXIS_CODE = "كود المحور"
C_AXIS = "المحور"
C_STAGE = "المرحلة"
C_OPENING = "السؤال الافتتاحي المقترح"
C_OUTPUT = "المخرج"

# The six roles, in the order they always appear inside a requirement.
ROLE_MAIN = "رئيسي"
ROLE_FOLLOWUP = "سؤال تابع"
ROLE_STATE = "الوضع الحالي"
ROLE_IMPL = "وسيلة التنفيذ"
ROLE_RULE = "حالة القاعدة"
ROLE_TARGET = "التوجه المستهدف"
STANDARD_ROLES = (ROLE_STATE, ROLE_IMPL, ROLE_RULE, ROLE_TARGET)

ROLE_KEYS = {
    ROLE_MAIN: "main",
    ROLE_FOLLOWUP: "followup",
    ROLE_STATE: "state",
    ROLE_IMPL: "impl",
    ROLE_RULE: "rule",
    ROLE_TARGET: "target",
}

QTYPES = {
    "اختيار واحد": "select",
    "اختيار متعدد": "multi",
    "جدول": "table",
    "رقم": "number",
}

# Options that are not a real answer. Mirrors azm.option.kind on the procurement side,
# so that "don't know" suppresses the follow-ups that depend on this question.
OPTION_KINDS = {
    "لا أعرف": "unknown",
    "غير مدرج": "not_listed",
    "لا يوجد": "none",
    "لا ينطبق": "na",
    "لا يوجد نظام": "none",
}

OPTION_SEP = "/"
CONDITION_RE = re.compile(r"يظهر\s+لو\s+(Q\d+)\s*=\s*(.+)", re.S)
CONDITION_OR = " أو "


def cell(row, headers, name):
    value = row[headers[name]]
    return "" if value is None else str(value).strip()


def read_sheet(path, title):
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    if title not in workbook.sheetnames:
        sys.exit(f"sheet {title!r} not found; got {workbook.sheetnames}")
    worksheet = workbook[title]
    rows = list(worksheet.iter_rows(values_only=True))
    headers = {str(h): i for i, h in enumerate(rows[0])}
    return headers, rows[1:]


def parse_options(raw):
    """Split the options cell. The sheet uses ' / ' consistently."""
    return [part.strip() for part in raw.split(OPTION_SEP) if part.strip()]


def parse_condition(raw, local_by_label):
    """Turn 'يظهر لو Q004 = a أو b' into the engine's condition grammar.

    Returns {"field": <trigger question key>, "in": [labels]} so that the existing
    engine.condition() can evaluate it unchanged.
    """
    if not raw:
        return None
    match = CONDITION_RE.search(raw)
    if not match:
        return {"__unparsed__": raw}
    trigger, tail = match.group(1), match.group(2)
    labels = [part.strip() for part in tail.split(CONDITION_OR) if part.strip()]
    return {"field": local_by_label.get(trigger, trigger), "in": labels}


def build(sheet_path):
    headers, rows = read_sheet(sheet_path, SHEET_QUESTIONS)
    o_headers, o_rows = read_sheet(sheet_path, SHEET_OPENING)
    previous = json.loads(CATALOGUE.read_text(encoding="utf-8"))
    shapes = {r["code"]: r for r in previous["requirements"]}

    modules, topics, requirements, questions = {}, {}, {}, []

    for row in rows:
        mcode = cell(row, headers, C_MODULE_CODE)
        tcode = cell(row, headers, C_TOPIC_CODE)
        rcode = cell(row, headers, C_REQ_CODE)
        qcode = cell(row, headers, C_Q_CODE)
        role = cell(row, headers, C_ROLE)
        key = qcode.rsplit("-", 1)[-1]

        modules.setdefault(mcode, {"code": mcode, "name": cell(row, headers, C_MODULE)})
        topics.setdefault(tcode, {"code": tcode, "name": cell(row, headers, C_TOPIC), "module": mcode})

        shape = shapes.get(rcode, {})
        requirement = requirements.setdefault(rcode, {
            "code": rcode,
            "uid": cell(row, headers, C_REQ_ID),
            "module": mcode,
            "topic": tcode,
            # Answer-shape metadata only the shipped catalogue has.
            "title": shape.get("title", cell(row, headers, C_TOPIC)),
            "profile": shape.get("profile", "L"),
            "scope": shape.get("scope", "ALL"),
            "shape": shape.get("shape", "table"),
            "columns": shape.get("columns", []),
            "branches": shape.get("branches", []),
            "adv": shape.get("adv", False),
            "pack": cell(row, headers, C_PACK),
            "owner": cell(row, headers, C_OWNER),
            "question": "",
        })

        options = parse_options(cell(row, headers, C_OPTIONS))
        question = {
            "code": f"{rcode}/{key}",
            "requirement": rcode,
            "key": key,
            "sequence": int(key.lstrip("Q") or 0),
            "role": ROLE_KEYS.get(role, role),
            "label": cell(row, headers, C_QUESTION),
            "qtype": QTYPES.get(cell(row, headers, C_FORMAT), "select"),
            "unit": cell(row, headers, C_UNIT),
            "columns": parse_options(cell(row, headers, C_COLUMNS)),
            "help": cell(row, headers, C_EXPECTED),
            "text_label": cell(row, headers, C_TEXTBOX),
            "response_type": cell(row, headers, C_RESPONSE),
            "condition_src": cell(row, headers, C_CONDITION),
            "options": [
                {"name": name, "sequence": i, "kind": OPTION_KINDS.get(name, "normal")}
                for i, name in enumerate(options, 1)
            ],
        }
        questions.append(question)
        # The main question's wording is the requirement's question.
        if role == ROLE_MAIN:
            requirement["question"] = question["label"]

    # Conditions reference a sibling by its local Qnnn code, which is already the key.
    for question in questions:
        question["condition"] = parse_condition(question.pop("condition_src"), {})

    opening = [{
        "node": cell(row, o_headers, C_NODE),
        "module": cell(row, o_headers, C_MODULE_CODE),
        "topic": cell(row, o_headers, C_TOPIC_CODE),
        "axis_code": cell(row, o_headers, C_AXIS_CODE),
        "axis": cell(row, o_headers, C_AXIS),
        "stage": cell(row, o_headers, C_STAGE),
        "question": cell(row, o_headers, C_OPENING),
        "output": cell(row, o_headers, C_OUTPUT),
    } for row in o_rows]

    # Sub-topics that only carry opening questions still need to exist.
    for item in opening:
        topics.setdefault(item["topic"], {
            "code": item["topic"],
            "name": cell(o_rows[0], o_headers, C_TOPIC),
            "module": item["module"],
        })
    for row in o_rows:
        tcode = cell(row, o_headers, C_TOPIC_CODE)
        topics[tcode]["name"] = cell(row, o_headers, C_TOPIC)

    # The 22 documented "عهد" detail questions hang off one module, whose code
    # changes with the new taxonomy (ADV -> FIN.ADV). Record it rather than
    # hard-coding it in the loader.
    adv_modules = {r["module"] for r in requirements.values() if r["adv"]}
    adv_module = sorted(adv_modules)[0] if len(adv_modules) == 1 else ""

    return {
        "version": 2,
        "source": f"AZM finance questionnaire · {pathlib.Path(sheet_path).name}",
        "adv_module": adv_module,
        "modules": sorted(modules.values(), key=lambda m: m["code"]),
        "topics": sorted(topics.values(), key=lambda t: t["code"]),
        "requirements": sorted(requirements.values(), key=lambda r: (r["module"], r["code"])),
        "questions": sorted(questions, key=lambda q: (q["requirement"], q["sequence"])),
        "opening_questions": opening,
        "adv_questions": previous.get("adv_questions", []),
    }


def validate(cat):
    """Every rule here exists because breaking it ships a wrong questionnaire."""
    errors = []
    by_req = defaultdict(list)
    for question in cat["questions"]:
        by_req[question["requirement"]].append(question)

    standard_scales = defaultdict(set)
    for question in cat["questions"]:
        if question["role"] in ("state", "impl", "rule", "target"):
            standard_scales[question["role"]].add(
                tuple(o["name"] for o in question["options"]))

    for role, scales in standard_scales.items():
        if len(scales) != 1:
            errors.append(f"standard role {role!r} has {len(scales)} different option sets")
    generic = {scale for scales in standard_scales.values() for scale in scales}

    for question in cat["questions"]:
        code, qtype, options = question["code"], question["qtype"], question["options"]
        names = [o["name"] for o in options]

        if qtype in ("select", "multi") and not options:
            errors.append(f"{code}: {qtype} question has no options")
        if qtype == "number":
            if options:
                errors.append(f"{code}: number question carries options")
            if not question["unit"]:
                errors.append(f"{code}: number question has no unit")
        if qtype == "table" and not question["columns"]:
            errors.append(f"{code}: table question has no columns")
        if len(names) != len(set(names)):
            dupes = [n for n, c in Counter(names).items() if c > 1]
            errors.append(f"{code}: duplicate option labels {dupes}")

        # The point of the whole exercise: a question must not answer with another
        # question's scale.
        if question["role"] in ("main", "followup") and tuple(names) in generic:
            errors.append(f"{code}: {question['role']} question carries a standard scale")

        condition = question["condition"]
        if not condition:
            continue
        if "__unparsed__" in condition:
            errors.append(f"{code}: cannot parse condition {condition['__unparsed__']!r}")
            continue
        siblings = {q["key"]: q for q in by_req[question["requirement"]]}
        trigger = siblings.get(condition["field"])
        if not trigger:
            errors.append(f"{code}: condition points at missing question {condition['field']}")
            continue
        if trigger["sequence"] >= question["sequence"]:
            errors.append(
                f"{code}: trigger {trigger['key']} does not sort before it; the engine "
                "only sees answers of earlier questions")
        if trigger["qtype"] == "multi":
            errors.append(
                f"{code}: trigger {trigger['key']} is multi-select; engine.condition() "
                "cannot match a list with 'in'")
        known = {o["name"] for o in trigger["options"]}
        unknown = [label for label in condition["in"] if label not in known]
        if unknown:
            errors.append(f"{code}: condition labels not offered by {trigger['key']}: {unknown}")

    codes = [q["code"] for q in cat["questions"]]
    if len(codes) != len(set(codes)):
        errors.append("duplicate question codes")
    topics = {t["code"] for t in cat["topics"]}
    for requirement in cat["requirements"]:
        if requirement["topic"] not in topics:
            errors.append(f"{requirement['code']}: unknown topic {requirement['topic']}")
    return errors


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("spreadsheet")
    parser.add_argument("-o", "--output", default=str(CATALOGUE))
    parser.add_argument("--check", action="store_true", help="validate only, write nothing")
    args = parser.parse_args()

    cat = build(args.spreadsheet)
    errors = validate(cat)

    print(f"modules            {len(cat['modules']):>6}")
    print(f"topics             {len(cat['topics']):>6}")
    print(f"requirements       {len(cat['requirements']):>6}")
    print(f"questions          {len(cat['questions']):>6}")
    print(f"options            {sum(len(q['options']) for q in cat['questions']):>6}")
    print(f"conditions         {sum(1 for q in cat['questions'] if q['condition']):>6}")
    print(f"opening questions  {len(cat['opening_questions']):>6}")

    if errors:
        print(f"\n{len(errors)} validation error(s):", file=sys.stderr)
        for error in errors[:40]:
            print(f"  - {error}", file=sys.stderr)
        sys.exit(1)
    print("\nvalidation passed")

    if args.check:
        return
    pathlib.Path(args.output).write_text(
        json.dumps(cat, ensure_ascii=False, indent=1), encoding="utf-8")
    print(f"written {args.output}")


if __name__ == "__main__":
    main()

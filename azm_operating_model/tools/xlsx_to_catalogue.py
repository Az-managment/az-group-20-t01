"""Build data/fin_catalogue.json from the finance questions workbook.

    python3 tools/xlsx_to_catalogue.py AZM_Finance_Questions_Extract.xlsx

Sheets used:
* «الأسئلة وطريقة الإجابة»  — one row per question (module, topic, requirement, role, answer form,
  unit, table columns, options, display condition, text box, V3 response type, pack, owner)
* «الأسئلة الافتتاحية»       — the model tree (modules / topics) and the draft opening questions per axis

Requirement applicability (scope) and answer profile are kept from the current catalogue,
matched on the previous requirement code (FIN-004, HRF-050 …).
"""
import json
import re
import sys
from pathlib import Path

import openpyxl

HERE = Path(__file__).resolve().parent.parent
OUT = HERE / "data" / "fin_catalogue.json"

ROLES = {"رئيسي": "main", "سؤال تابع": "follow", "الوضع الحالي": "status", "وسيلة التنفيذ": "impl",
         "حالة القاعدة": "rule", "التوجه المستهدف": "target"}
QTYPES = {"اختيار واحد": "select", "اختيار متعدد": "multi", "جدول": "table", "رقم": "number"}
KINDS = {"لا أعرف": "unknown", "غير مدرج": "not_listed", "لا يوجد": "none", "لا يوجد نظام": "none",
         "لا تتم": "none", "لا ينطبق": "none"}
STAGES = {"حصر الوضع الحالي": "current", "تصميم لاحق بعد المراجعة": "design"}


def clean(value):
    return re.sub(r"\s+", " ", str(value)).strip() if value is not None else ""


def parse_options(text):
    return [{"name": o, "kind": KINDS.get(o, "")} for o in (clean(x) for x in clean(text).split(" / ")) if o]


PEOPLE = ["المدير المالي", "رئيس الحسابات", "محاسب", "محاسب الموردين", "محاسب العملاء", "مسؤول الموازنة",
         "مسؤول الخزينة", "مسؤول الرواتب", "الموارد البشرية", "المشتريات", "مدير الإدارة الطالبة",
         "الشؤون القانونية", "المراجعة الداخلية", "مدير الكيان", "المدير المالي للمجموعة", "الرئيس التنفيذي",
         "مجلس الإدارة", "تقنية المعلومات"]
ROLE_COLUMNS = {"يعتمد", "المسؤول", "يراجع", "المنفذ", "الدور", "المعتمد", "يُنشئ", "يطلب", "الضابط البديل"}
NUMBER_COLUMNS = {"النسبة %", "القيمة أو النسبة", "الحد الأقصى", "القيمة", "حد القيمة", "بعد (يوم)", "نسبة الاسترداد %",
                  "من مبلغ", "إلى مبلغ", "حد الصرف للمرة", "البدل اليومي", "العمر (سنة)", "القيمة المتبقية %",
                  "نسبة الموظف %", "نسبة صاحب العمل %", "حد الساعات الشهري", "الحد الأقصى %", "نسبة الأجر المدفوع %",
                  "يوم العمل", "الأولوية"}
DATE_COLUMNS = {"تاريخ الاستحقاق", "ينتهي في"}
BASES = ["نسبة من الأساسي", "نسبة من الإجمالي", "مبلغ ثابت", "حسب الشرائح", "فعلي بمستندات", "حسب الساعات",
         "حسب التكلفة الفعلية", "تقدير اكتواري", "حسب لائحة معتمدة"]
TREATMENTS = ["قيد تسوية", "إثبات مصروف", "إثبات إيراد", "تحميل على مركز تكلفة أو مشروع", "خصم من المستحقات",
              "استرداد نقدي", "استهلاك على الفترة", "إثبات التزام", "معالجة يدوية خارج النظام"]
ACCOUNTS = ["مصروف الرواتب والأجور", "مخصص الإجازات", "مخصص مكافأة نهاية الخدمة", "مخصص المكافآت", "مستحقات العاملين",
            "مصروفات مستحقة", "النقدية والبنوك", "حساب آخر"]
LIST_COLUMNS = {
    "المستند": ["طلب معتمد", "أمر شراء", "عقد", "فاتورة", "محضر استلام", "مستخلص", "سند صرف", "سند قبض", "قيد يومية",
                "كشف حساب بنكي", "مذكرة داخلية", "نموذج إلكتروني في النظام"],
    "المستند المطلوب": ["مستند أصلي معتمد", "فاتورة", "عقد", "كشف حساب", "محضر", "مذكرة تسوية معتمدة", "جدول احتساب",
                        "موافقة بريد إلكتروني", "تقرير من النظام"],
    "الدليل": ["اعتماد في النظام", "مستند ورقي موقع", "بريد إلكتروني", "تقرير من النظام", "محضر", "كشف مطابقة"],
    "التقرير": ["قائمة الدخل", "الميزانية العمومية", "التدفقات النقدية", "ميزان المراجعة", "الموازنة مقابل الفعلي",
                "أعمار الديون", "تقرير السيولة", "تكلفة الرواتب", "تقرير المشروعات", "تقرير المانح", "تقرير إداري مخصص"],
    "المستفيد": ["مجلس الإدارة", "الرئيس التنفيذي", "المدير المالي للمجموعة", "المدير المالي", "مدير الكيان", "الإدارات",
                 "المانح", "البنوك", "الجهات الرقابية", "المراجع الخارجي"],
    "الجهة": ["مصلحة الضرائب", "الهيئة القومية للتأمين الاجتماعي", "الجهاز المركزي للمحاسبات", "وزارة التعليم العالي",
              "المجلس الأعلى للجامعات", "الهيئة العامة للرقابة المالية", "البنك المركزي", "جهة أخرى"],
    "أساس الحساب": BASES, "الأساس": BASES, "طريقة الحساب": BASES, "طريقة القياس": BASES + ["دراسة اكتوارية مستقلة"],
    "أساس التسعير": ["سعر ثابت للخدمة", "حسب التكلفة الفعلية", "حسب المرحلة الدراسية", "حسب الاستخدام", "سعر السوق"],
    "المعالجة": TREATMENTS, "المعالجة الحالية": TREATMENTS, "التسجيل": TREATMENTS,
    "الاسترداد أو الاستهلاك": ["استهلاك من المستخلصات", "استرداد عند التسليم", "استرداد نهاية الضمان", "خصم من الدفعات",
                               "استهلاك على مدة العقد"],
    "الأثر المالي": ["منخفض", "متوسط", "مرتفع"],
    "مدين": ACCOUNTS, "دائن": ACCOUNTS, "حساب المصروف": ACCOUNTS,
    "البنك": ["البنك الأهلي المصري", "بنك مصر", "البنك التجاري الدولي", "بنك القاهرة", "بنك الإسكندرية",
              "بنك قطر الوطني الأهلي", "بنك أبوظبي الأول", "بنك آخر"],
    "الحساب البنكي": ["الحساب الجاري الرئيسي", "حساب الرواتب", "حساب التحصيل", "حساب بالعملة الأجنبية", "حساب ودائع"],
    "الضمان": ["رهن عقاري", "رهن أصول", "كفالة الشركة الأم", "ودائع مجمدة", "تنازل عن مستحقات", "بدون ضمان"],
    "الضريبة": ["ضريبة القيمة المضافة", "ضريبة الخصم والإضافة", "ضريبة الدخل", "ضريبة الدمغة", "ضريبة كسب العمل"],
    "النظام": ["نظام الطلاب", "نظام الموارد البشرية", "نظام الحضور والانصراف", "نظام نقاط البيع", "نظام المخازن",
               "نظام المشروعات", "نظام العقارات", "Excel", "نظام آخر"],
    "المعاملات": ["رسوم وتحصيلات", "مشتريات وفواتير", "رواتب ومستحقات", "حركات مخزنية", "تكاليف مشروعات",
                  "إيجارات وأقساط", "معاملات أخرى"],
    "الفئة": ["أراضٍ", "مبانٍ", "أثاث ومفروشات", "أجهزة حاسب", "سيارات", "معدات وآلات", "تجهيزات معامل", "برمجيات"],
    "القسط": ["القسط الأول", "القسط الثاني", "القسط الثالث", "القسط الرابع", "القسط الخامس", "القسط السادس"],
    "الدرجة الوظيفية": ["الإدارة العليا", "مدير إدارة", "رئيس قسم", "أخصائي", "موظف", "عضو هيئة تدريس", "عمالة"],
    "الفئات المستحقة": ["كل العاملين", "الإداريون", "أعضاء هيئة التدريس", "الإدارة العليا", "العمالة", "فئات محددة"],
    "مدة الاحتفاظ": ["سنة", "3 سنوات", "5 سنوات", "10 سنوات", "دائم"],
    "السريان المعتاد": ["أول السنة المالية", "أول الشهر التالي", "من تاريخ القرار", "بأثر رجعي"],
    "ينتهي في أو بانتهاء": ["تاريخ محدد", "انتهاء العقد", "انتهاء المشروع", "نهاية السنة المالية", "عند الإلغاء"],
    "المهام المجمعة": ["الإنشاء والاعتماد", "الإنشاء والمراجعة", "المراجعة والاعتماد", "الإدخال والدفع",
                       "الدفع والتسوية البنكية", "إنشاء المورد والدفع له", "الاستلام والفوترة"],
    "الشريحة": ["الشريحة الأولى", "الشريحة الثانية", "الشريحة الثالثة", "الشريحة الرابعة", "الشريحة الخامسة"],
}


def type_column(col):
    """Give every table column the most specific input: a list wherever possible."""
    label = col["label"]
    if col.get("options"):
        col["type"] = "list"
    elif label in ROLE_COLUMNS:
        col.update(type="list", options=PEOPLE + ["لا يوجد"])
    elif label in NUMBER_COLUMNS:
        col["type"] = "number"
    elif label in DATE_COLUMNS:
        col["type"] = "date"
    elif label in LIST_COLUMNS:
        col.update(type="list", options=LIST_COLUMNS[label])
    else:
        col["type"] = "text"
    return col


def parse_columns(text):
    cols = []
    for part in clean(text).split(" / "):
        m = re.match(r"^(.*?)\s*\((.*)\)\s*$", part)
        if m and "،" in m.group(2):
            cols.append({"label": m.group(1).strip(), "options": [o.strip() for o in m.group(2).split("،") if o.strip()]})
        elif part:
            cols.append({"label": part})
    return [type_column(c) for c in cols]


def parse_condition(text, questions):
    """«يظهر لو Q003 = A أو B» → {"q": "Q003", "in": [A, B]} (option names may contain «أو»)."""
    text = clean(text)
    if not text:
        return None
    m = re.match(r"^يظهر لو (Q\d+) = (.+)$", text)
    if not m:
        raise ValueError(f"unknown condition: {text}")
    ref, rest = m.groups()
    names = [o["name"] for o in questions[ref]["options"]]
    values, buf = [], []
    for token in rest.split(" أو "):
        buf.append(token)
        candidate = " أو ".join(buf)
        if candidate in names:
            values.append(candidate)
            buf = []
    if buf or not values:
        raise ValueError(f"condition values not in {ref} options: {text}")
    return {"q": ref, "in": values}


def main(path):
    wb = openpyxl.load_workbook(path, read_only=True)
    old = json.loads(OUT.read_text()) if OUT.exists() else {}
    old_reqs = {r["code"]: r for r in old.get("requirements", [])}

    # ---------------------------------------------------------------- tree + opening questions
    modules, topics = {}, {}
    for row in wb["الأسئلة الافتتاحية"].iter_rows(min_row=2, values_only=True):
        node, mcode, mname, tcode, tname, acode, aname, stage, question, output = (clean(c) for c in row[:10])
        if not node:
            continue
        modules.setdefault(mcode, {"code": mcode.split(".", 1)[1], "full_code": mcode, "name": mname})
        topic = topics.setdefault(tcode, {"code": tcode, "module": mcode.split(".", 1)[1], "name": tname, "prompts": []})
        topic["prompts"].append({"code": node, "axis": acode, "axis_name": aname, "stage": STAGES.get(stage, "current"),
                                 "question": question, "output": output})

    # ---------------------------------------------------------------- questions
    reqs = {}
    for row in wb["الأسئلة وطريقة الإجابة"].iter_rows(min_row=2, values_only=True):
        (mcode, mname, tcode, tname, req_id, prev, qcode, role, label, hint, form, unit, cols, options,
         cond, note, rtype, pack, owner) = (clean(c) for c in row[:19])
        if not qcode:
            continue
        modules.setdefault(mcode, {"code": mcode.split(".", 1)[1], "full_code": mcode, "name": mname})
        topics.setdefault(tcode, {"code": tcode, "module": mcode.split(".", 1)[1], "name": tname, "prompts": []})
        req = reqs.setdefault(prev, {
            "code": prev, "req_id": req_id, "module": mcode.split(".", 1)[1], "topic": tcode,
            "response_type": rtype, "pack": pack, "owner": owner, "questions": [], "_by_key": {},
        })
        key = qcode.rsplit("-", 1)[1]
        q = {
            "code": qcode, "key": key, "role": ROLES[role], "label": label, "hint": hint,
            "qtype": QTYPES[form], "unit": unit, "columns": parse_columns(cols) if form == "جدول" else [],
            "options": parse_options(options), "condition_text": cond, "note_label": note,
        }
        req["_by_key"][key] = q
        req["questions"].append(q)

    for req in reqs.values():
        for q in req["questions"]:
            q["condition"] = parse_condition(q.pop("condition_text"), req["_by_key"])
        main_q = next(q for q in req["questions"] if q["role"] == "main")
        prev = old_reqs.get(req["code"], {})
        req.update({
            "title": prev.get("title") or main_q["label"],
            "question": main_q["label"],
            "profile": prev.get("profile", "P"),
            "scope": prev.get("scope", "ALL"),
        })
        del req["_by_key"]

    order = {code: i for i, code in enumerate(topics)}
    requirements = sorted(reqs.values(), key=lambda r: (order[r["topic"]], r["req_id"]))
    data = {
        "source": "AZM_Finance_Questions_Extract.xlsx",
        "modules": list(modules.values()),
        "topics": list(topics.values()),
        "requirements": requirements,
    }
    OUT.write_text(json.dumps(data, ensure_ascii=False, separators=(",", ":")))
    nq = sum(len(r["questions"]) for r in requirements)
    print(f"{len(modules)} modules · {len(topics)} topics · {len(requirements)} requirements · {nq} questions"
          f" · {sum(len(t['prompts']) for t in topics.values())} opening questions → {OUT}")


if __name__ == "__main__":
    main(sys.argv[1])
